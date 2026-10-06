import json, os, cmo_lib as L, maps
L.HC=os.path.expanduser('~/w/out5/res')
AM=json.load(open('auto_maps5.json'))
maps.W.update(AM['W']); maps.S.update(AM['S'])
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
B=json.load(open('built5.json')); D=L.load_cmo(); H=L.load_hc()
names={}
for k in ['c1','d1','e1','f1','g1','g2','j1','m1']:
    for r in H[k][1]: names[int(r['id'])]=r['classname']
for k in ['c1','d1','f1','g2','j1','m1']:
    for r in B['NEW'][k]: names[int(r['id'])]=r['classname']
wb=Workbook(); F=Font(name='Arial',size=10); HB=Font(name='Arial',size=10,bold=True,color='FFFFFF'); FILL=PatternFill('solid',fgColor='1F3864')
def sheet(ws,hdr,rows,widths):
    ws.append(hdr)
    for r in rows: ws.append(r)
    for c in ws[1]: c.font=HB; c.fill=FILL
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=F; c.alignment=Alignment(vertical='top',wrap_text=False)
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
ws=wb.active; ws.title='Platforms'
rows=[]
for l in B['LOG']:
    if l[0] in('ship','sub','aircraft'): rows.append([l[0],l[1],l[2],int(l[3]),l[5],l[4]])
sheet(ws,['Kind','HCDB name','Annex','HCDB ID','CMO source','Build notes'],rows,[10,28,7,9,60,110])
ws=wb.create_sheet('New weapons & sensors')
rows=[[l[0],l[1],l[2],int(l[3]),l[5],l[4]] for l in B['LOG'] if l[0] in('weapon','sensor')]
sheet(ws,['Kind','HCDB name','Annex','HCDB ID','CMO source','How built'],rows,[10,28,7,9,45,50])
ws=wb.create_sheet('Weapon map')
rows=[]
for c,v in sorted(maps.W.items(),key=lambda x:D['Weapon'][x[0]]['Name'] if x[0] in D['Weapon'] else ''):
    if c not in D['Weapon']: continue
    if isinstance(v,int): t,h='existing',f'{v} {names.get(v,"")}'
    elif isinstance(v,str): t,h=v.lower(),''
    elif v[0]=='DIV': t,h='existing (qty / %d)'%v[2],f'{v[1]} {names.get(v[1],"")}'
    else: t,h='new (cloned from %s)'%names.get(v[1],v[1]),v[2]
    rows.append([int(c),D['Weapon'][c]['Name'],t,h])
sheet(ws,['CMO weapon ID','CMO name','Treatment','HCDB item'],rows,[12,60,18,45])
ws=wb.create_sheet('Sensor map')
rows=[]
for c,v in maps.S.items():
    if c not in D['Sensor']: continue
    rows.append([int(c),D['Sensor'][c]['Name'],'existing' if isinstance(v,int) else ('skip' if isinstance(v,str) else 'new'),f'{v} {names.get(v,"")}' if isinstance(v,int) else (v if isinstance(v,str) else v[2])])
sheet(ws,['CMO sensor ID','CMO name','Treatment','HCDB item'],rows,[12,50,12,40])
ws=wb.create_sheet('Notes')
for t in ['CMO DB3000 -> HCDB conversion, 5 Oct 2026: US batch (89 platforms) plus 1158 platforms (409 ships, 54 subs, 695 aircraft): 2015+ equipment for every country in CMO, plus older (1980-2014 service) ships and subs for the 44 countries Harpoon Classic did not have (Argentina, Brazil, Chile, Peru, Mexico, Cuba and others).',
          'Each new platform was cloned from the closest existing HCDB entry (see Build notes) and its weapons, sensors, mounts, magazines and loadouts rebuilt from CMO.',
          'New IDs use the lower half of each annex (the range HCDB-250724 itself uses for additions). Existing rows were not changed.',
          'Artwork reuses the analog platform picture. Descriptions are auto-generated and can be edited in annex_strings.csv.',
          'Radars/sonars only are carried over; ESM/IR/MAD become unit flags where applicable. Ship chaff/decoys use the generic Naval decoy/ECM mounts. Towed torpedo decoys are omitted.',
          'Repeated launchers were consolidated (e.g. 24 SA-N-9 modules on Kuznetsov = 8 mounts x 24 rounds); names transliterated to Latin. Identical aircraft loadouts are shared between aircraft to save database space. Loadouts chosen for mission variety, then to cover every carriable weapon (up to 10 per aircraft); US batch aircraft loadouts regenerated the same way.',
          'CMO data is copyright WarfareSims/Matrix Games - personal use only.']:
    wb['Notes'].append([t])
wb['Notes'].column_dimensions['A'].width=150
for r in wb['Notes'].iter_rows():
    for c in r: c.font=F
import os; os.makedirs('/mnt/user-data/outputs',exist_ok=True)
wb.save('/mnt/user-data/outputs/CMO_to_HCDB_review_v4.xlsx')
