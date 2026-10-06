import json, csv, os, re, collections, sys
sys.path.insert(0, os.path.expanduser('~/w'))
import cmo_lib as L
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
W = os.path.expanduser('~/w')
OUT = f'{W}/out5/res'
def rd(f):
    rows = list(csv.reader(open(f'{OUT}/annex_{f}.csv', encoding='latin-1'), skipinitialspace=True))
    h = [x.strip() for x in rows[0]]
    return {int(r[0]): dict(zip(h, r)) for r in rows[1:] if r}
T = {k: rd(k) for k in ['a1','a2','b1','c1','d1','f1','g2','j1','m1','y1','z1','z2']}
for line in open(f'{OUT}/annex_bs_strings.csv', encoding='latin-1'):
    if line.startswith('"CountryName"'):
        CN = next(csv.reader([line]))[3:]
B1 = json.load(open(f'{W}/built.json')); B5 = json.load(open(f'{W}/built5.json'))
P1 = json.load(open(f'{W}/platforms.json')); P5 = json.load(open(f'{W}/platforms5.json'))
batch = {}
for n, f in [('2 – China/Russia/Koreas/Japan/NATO', 'selection2.json'), ('3 – India/Israel/Iran/Australia/Taiwan + gap fill', 'selection3.json'),
             ('4 – rest of world (2015+)', 'selection4.json'), ('5 – older ships/subs, new countries', 'selection5.json')]:
    for s in json.load(open(f'{W}/{f}')):
        batch[s['kind'] + '|' + s['country'] + '|' + s['key']] = n
plat = []
def add(logs, P, bname_fn):
    ks = list(P)
    for kind, annex in [('ship', 'a1'), ('sub', 'a2'), ('aircraft', 'b1')]:
        kk = [k for k in ks if (P[k]['kind'] == {'ship':'ship','sub':'sub','aircraft':'air'}[kind])]
        ll = [x for x in logs if x[0] == kind]
        for k, x in zip(kk, ll):
            r = T[annex].get(int(x[3]))
            if not r: continue
            p = P[k]
            m = re.search(r'analog (.*?);', x[4])
            plat.append(dict(batch=bname_fn(k), kind={'ship':'Ship','sub':'Submarine','aircraft':'Aircraft'}[kind], annex=annex, id=int(x[3]),
                 name=r['classname'], country=CN[int(r['country'])], year=p['year'], cmo=p['cmoname'], analog=m.group(1) if m else '',
                 disp=int(r['displacement']) if annex != 'b1' else None, dp=int(r['origdp']) if annex != 'b1' else None,
                 speed=int(r['maxspeed']) if annex == 'a1' else (int(r['maxsubspeed']) if annex == 'a2' else None),
                 mounts=len(L.arr(r['mount_array'])) if annex != 'b1' else None,
                 sensors=len(L.arr(r['sensor_array'])), loadouts=len(L.arr(r['Loadout_array'])) if annex == 'b1' else None))
add(B1['LOG'], P1, lambda k: '1 – United States')
add(B5['LOG'], P5, lambda k: batch.get(k, '?'))
# weapons & sensors
ws = []; seen = set()
for B in (B1, B5):
    for x in B['LOG']:
        if x[0] in ('weapon', 'sensor'):
            i = int(x[3])
            if (x[2], i) in seen or i not in T.get(x[2], {}): continue
            seen.add((x[2], i))
            r = T[x[2]][i]
            ws.append(dict(kind={'c1':'Gun','d1':'Missile','f1':'Torpedo/mine','g2':'Pod/ECM','j1':'Radar','m1':'Sonar'}.get(x[2], x[2]),
                           annex=x[2], id=i, name=r['classname'], cmo=x[5], how=x[4]))
# modifications to existing data
mods = []
for i, n in B5['NEW_COUNTRIES']:
    mods.append(('Country list (bs_strings CountryName)', i, n, 'Added country name in empty slot'))
for bid, lo in B5['US_LO'].items():
    r = T['b1'][int(bid)]
    mods.append(('Aircraft loadouts (b1 Loadout_array)', int(bid), r['classname'], f'Loadouts regenerated for weapon coverage ({len(L.arr(lo))} loadouts); old US-batch y1 rows 53248-53589 removed'))
json.dump({'plat': plat, 'ws': ws, 'mods': mods}, open(f'{W}/pub/catalog.json', 'w'))
# ---------- workbook
F = Font(name='Arial', size=10); HB = Font(name='Arial', size=10, bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='1F3864')
TF = Font(name='Arial', size=14, bold=True)
wb = Workbook()
def sheet(ws_, hdr, rows, widths):
    ws_.append(hdr)
    for r in rows: ws_.append(r)
    for c in ws_[1]: c.font = HB; c.fill = FILL; c.alignment = Alignment(vertical='center')
    for row in ws_.iter_rows(min_row=2):
        for c in row: c.font = F
    for i, w in enumerate(widths): ws_.column_dimensions[get_column_letter(i + 1)].width = w
    ws_.freeze_panes = 'A2'; ws_.auto_filter.ref = ws_.dimensions
s0 = wb.active; s0.title = 'Summary'
wsP = wb.create_sheet('Platforms')
plat.sort(key=lambda p: (p['batch'], p['kind'], p['country'], p['name']))
sheet(wsP, ['Batch', 'Type', 'Annex', 'HCDB ID', 'HCDB name', 'Country', 'Year', 'CMO source platform', 'Cloned from (analog)',
            'Displacement (t)', 'Damage points', 'Max speed (kt)', 'Mounts', 'Sensors', 'Loadouts'],
      [[p['batch'], p['kind'], p['annex'], p['id'], p['name'], p['country'], p['year'], p['cmo'], p['analog'], p['disp'], p['dp'], p['speed'], p['mounts'], p['sensors'], p['loadouts']] for p in plat],
      [34, 11, 7, 9, 30, 16, 7, 50, 30, 12, 11, 11, 9, 9, 9])
wsW = wb.create_sheet('Weapons & Sensors')
ws.sort(key=lambda w: (w['kind'], w['name']))
sheet(wsW, ['Type', 'Annex', 'HCDB ID', 'HCDB name', 'CMO source', 'How built'], [[w['kind'], w['annex'], w['id'], w['name'], w['cmo'], w['how']] for w in ws], [13, 7, 9, 30, 55, 45])
wsM = wb.create_sheet('Modified existing')
sheet(wsM, ['What changed', 'ID / index', 'Name', 'Change'], [list(m) for m in mods], [38, 11, 34, 80])
wsC = wb.create_sheet('Countries')
cts = sorted(set(p['country'] for p in plat))
sheet(wsC, ['Country', 'Ships', 'Submarines', 'Aircraft', 'Total'], [], [22, 9, 11, 9, 9])
n = len(plat) + 1
for i, c in enumerate(cts, start=2):
    wsC.cell(i, 1, c).font = F
    for j, k in enumerate(['Ship', 'Submarine', 'Aircraft'], start=2):
        wsC.cell(i, j, f"=COUNTIFS(Platforms!$F$2:$F${n},$A{i},Platforms!$B$2:$B${n},\"{k}\")").font = F
    wsC.cell(i, 5, f'=SUM(B{i}:D{i})').font = F
wsC.auto_filter.ref = f'A1:E{len(cts)+1}'
# summary
s0['A1'] = 'HCDB-250724 + CMO 2026 additions'; s0['A1'].font = TF
s0['A2'] = 'Platform, weapon and sensor additions to the Harpoon Classic / HUCE commondb.res, converted from the Command: Modern Operations database.'; s0['A2'].font = F
rows = [('Ships', "=COUNTIF(Platforms!$B$2:$B${n},\"Ship\")"), ('Submarines', f"=COUNTIF(Platforms!$B$2:$B${n},\"Submarine\")"),
        ('Aircraft', f"=COUNTIF(Platforms!$B$2:$B${n},\"Aircraft\")")]
s0['A4'] = 'New platforms'; s0['A4'].font = Font(name='Arial', size=11, bold=True)
r0 = 5
for i, (a, f) in enumerate(rows):
    s0.cell(r0 + i, 1, a).font = F; s0.cell(r0 + i, 2, f.replace('${n}', str(n))).font = F
s0.cell(r0 + 3, 1, 'Total platforms').font = Font(name='Arial', size=10, bold=True); s0.cell(r0 + 3, 2, f'=SUM(B{r0}:B{r0+2})').font = Font(name='Arial', size=10, bold=True)
m = len(ws) + 1
s0['A10'] = 'New weapons & sensors'; s0['A10'].font = Font(name='Arial', size=11, bold=True)
for i, k in enumerate(['Missile', 'Gun', 'Torpedo/mine', 'Pod/ECM', 'Radar', 'Sonar']):
    s0.cell(11 + i, 1, k).font = F; s0.cell(11 + i, 2, f'=COUNTIF(\'Weapons & Sensors\'!$A$2:$A${m},A{11+i})').font = F
s0['A18'] = 'By batch'; s0['A18'].font = Font(name='Arial', size=11, bold=True)
bs = sorted(set(p['batch'] for p in plat))
for i, b in enumerate(bs):
    s0.cell(19 + i, 1, b).font = F; s0.cell(19 + i, 2, f'=COUNTIF(Platforms!$A$2:$A${n},A{19+i})').font = F
k = 19 + len(bs) + 1
s0.cell(k, 1, 'Countries with new platforms').font = F; s0.cell(k, 2, f'=COUNTA(Countries!A2:A{len(cts)+1})').font = F
s0.cell(k + 1, 1, 'Existing rows modified').font = F; s0.cell(k + 1, 2, f"=COUNTA('Modified existing'!A2:A{len(mods)+1})").font = F
notes = ['Notes',
 'Base database: HCDB-250724 (HarpGamer). Every new item uses free IDs in the lower half of each annex block; no existing platform, weapon or sensor row was changed except those listed on "Modified existing".',
 'Each new platform is cloned from the closest existing HCDB entry (picture, flags, base handling) and then rebuilt from CMO: mounts, magazines, sensors, loadouts, displacement, damage points, speed.',
 'Damage points use the HCDB formula (T/20, T/30+9, T/50+76, T/60+116 by displacement, with HCDB multipliers).',
 'Aircraft loadouts: mission variety first, then greedy coverage of every weapon the aircraft can carry, max 10 per aircraft; identical loadouts are shared.',
 'Only radars and sonars are carried over; ESM/IR/MAD become unit flags. Cargo, troops, recon pods, lasers and towed torpedo decoys are not represented.',
 'Source data: Command: Modern Operations DB3000 (WarfareSims / Matrix Games). Derived values; for personal use unless the rights holders agree otherwise.']
for i, t in enumerate(notes):
    c = s0.cell(k + 3 + i, 1, t); c.font = Font(name='Arial', size=11 if i == 0 else 10, bold=(i == 0))
s0.column_dimensions['A'].width = 52; s0.column_dimensions['B'].width = 12
out = '/mnt/user-data/outputs/HCDB_CMO_additions_catalog.xlsx'
wb.save(out); print(out, len(plat), len(ws), len(mods), len(cts))
